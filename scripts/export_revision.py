#!/usr/bin/env python3
"""Exporta la base de conocimiento a una hoja de cálculo de revisión clínica (xlsx).

Uso:
    python3 scripts/export_revision.py                 # → data/revision/revision_<fecha>.xlsx
    python3 scripts/export_revision.py salida.xlsx

Hojas: Instrucciones · Reglas · Umbrales · Reglas combinadas · Interacciones · Clases.
En cada hoja el revisor rellena solo las columnas ESTADO (OK / CAMBIAR / ELIMINAR), TEXTO NUEVO y COMENTARIO.
La columna oculta "ref" identifica cada fila para que scripts/import_revision.py pueda aplicar los cambios.
"""
import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
KB = json.loads((ROOT / "app/data/kb.json").read_text(encoding="utf-8"))
DDI = json.loads((ROOT / "app/data/ddi_curated.json").read_text(encoding="utf-8"))
try:
    FICHAS = json.loads((ROOT / "app/data/fichas.json").read_text(encoding="utf-8"))
except Exception:  # noqa: BLE001
    FICHAS = {}
AMBITOS = {"general": "Antes de empezar", "aerobico": "Aeróbico", "fuerza": "Fuerza", "equilibrio": "Equilibrio"}
NIVEL = {"alerta": "ALERTA", "precaucion": "precaución", "info": "info"}

MORADO = "6B35C4"
GRIS = "F2F0F7"
CAB = Font(bold=True, color="FFFFFF")
FILL_CAB = PatternFill("solid", fgColor=MORADO)
FILL_REV = PatternFill("solid", fgColor="FFF4D6")  # columnas que rellena el revisor
WRAP = Alignment(wrap_text=True, vertical="top")

porId = {p["id"]: (c, p) for c in KB["clases"] for p in c["principios"]}


def nombre_ref(k):
    if k.startswith("clase:"):
        c = next((c for c in KB["clases"] if c["id"] == k[6:]), None)
        return f"[grupo] {c['nombre'] if c else k[6:]}"
    return porId[k][1]["nombre"] if k in porId else k


def condicion(si):
    if not si:
        return ""
    partes = []
    if si.get("dosis"):
        partes.append(f"dosis {si['dosis']}")
    if si.get("principio"):
        partes.append("solo " + (porId[si["principio"]][1]["nombre"] if si["principio"] in porId else si["principio"]))
    if si.get("edad_min"):
        partes.append(f"≥ {si['edad_min']} años")
    if si.get("condicion"):
        partes.append("si " + si["condicion"])
    if si.get("sin_condicion"):
        partes.append("si no " + si["sin_condicion"])
    if si.get("caidas"):
        partes.append("con caídas")
    if si.get("senal"):
        partes.append(f"≥ {si.get('min', 2)} fármacos con señal '{si['senal']}'")
    if si.get("todas"):
        partes.append("todas: " + ", ".join(si["todas"]))
    return "; ".join(partes)


def hoja(wb, titulo, columnas, filas, anchos, rev_desde, ref_col=True):
    ws = wb.create_sheet(titulo)
    cols = columnas + (["ref"] if ref_col else [])
    ws.append(cols)
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = CAB
        cell.fill = FILL_CAB
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for f in filas:
        ws.append(f)
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    if ref_col:
        ws.column_dimensions[get_column_letter(len(cols))].hidden = True
    for row in ws.iter_rows(min_row=2, max_col=len(cols)):
        for cell in row:
            cell.alignment = WRAP
        for cell in row[rev_desde - 1:len(columnas)]:
            cell.fill = FILL_REV
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{ws.max_row}"
    dv = DataValidation(type="list", formula1='"OK,CAMBIAR,ELIMINAR"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{get_column_letter(rev_desde)}2:{get_column_letter(rev_desde)}{max(ws.max_row, 2)}")
    return ws


def main(out: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Instrucciones"
    txt = [
        ["Revisión clínica · Medicación y Ejercicio", f"Base de conocimiento v{KB['version']} · exportada el {date.today():%d/%m/%Y}"],
        [""],
        ["Cómo revisar", "En cada hoja, las columnas amarillas son las suyas. Deje ESTADO en blanco o ponga OK si la fila está bien."],
        ["", "CAMBIAR + TEXTO NUEVO (o NIVEL/UMBRAL NUEVO): se sustituye por lo que escriba. Puede copiar el texto original, pegarlo y retocarlo."],
        ["", "ELIMINAR: la regla desaparece de la app."],
        ["", "COMENTARIO: cualquier nota (duda, fuente, 'esto va en otro grupo'…). Se lee pero no se aplica automáticamente."],
        ["", "No borre ni reordene filas ni toque la columna oculta 'ref': es la que permite aplicar sus cambios."],
        ["Niveles", "ALERTA = puede impedir o condicionar seriamente el ejercicio · precaución = adaptar · info = solo orientación."],
        ["Ámbitos", "Antes de empezar (general) · Aeróbico · Fuerza · Equilibrio."],
        ["Condición", "Cuándo se muestra la regla: 'dosis alta' (por encima del umbral de la hoja Umbrales), edad, situación clínica, solo un principio del grupo…"],
        ["Fuente", "Si en COMENTARIO escribe la fuente al principio (p. ej. «FT 4.4», «ACSM ap. A», «Lexicomp»), se guarda junto a la regla y se muestra en la aplicación."],
        ["Al terminar", "Guarde el archivo (si es Google Sheets: Archivo → Descargar → Excel) y envíelo; scripts/import_revision.py aplica los cambios y regenera la base."],
        ["Hojas", "Reglas (612) · Umbrales de dosis · Reglas combinadas (por varios fármacos o situación) · Interacciones revisadas · Clases (resumen y señales)."],
    ]
    for r in txt:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=14, color=MORADO)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 120
    for row in ws.iter_rows(min_row=3):
        row[0].font = Font(bold=True)
        row[1].alignment = WRAP

    # --- Reglas por clase y ámbito
    filas = []
    for c in KB["clases"]:
        for amb, nom in AMBITOS.items():
            for i, r in enumerate(c.get(amb, [])):
                fts = [FICHAS[p["id"]] for p in c["principios"] if p["id"] in FICHAS][:1]
                filas.append([c["nombre"], nom, NIVEL.get(r["nivel"], r["nivel"]), condicion(r.get("si")), r["texto"], "Sí" if r.get("revisado") else "", fts[0] if fts else "", "", "", "", f"regla|{c['id']}|{amb}|{i}"])
    hoja(wb, "Reglas", ["Grupo", "Ámbito", "Nivel", "Condición", "Texto actual", "Revisada", "Ficha técnica (1.er principio)", "ESTADO", "TEXTO NUEVO", "COMENTARIO"], filas,
         [30, 14, 11, 24, 80, 9, 24, 11, 60, 30], rev_desde=8)

    # --- Umbrales y principios
    filas = []
    for c in KB["clases"]:
        for p in c["principios"]:
            if p["id"].startswith("atc:"):
                continue
            filas.append([c["nombre"], p["nombre"], ", ".join(p.get("marcas", [])[:6]), p.get("unidad", "mg"),
                          p.get("umbral_mg", {}).get("alta"), p.get("nota", ""), FICHAS.get(p["id"], ""), "", "", "", f"principio|{c['id']}|{p['id']}"])
    hoja(wb, "Umbrales", ["Grupo", "Principio activo", "Marcas", "Unidad", "Dosis alta a partir de (por día)", "Nota", "Ficha técnica", "ESTADO", "UMBRAL NUEVO", "COMENTARIO"], filas,
         [30, 34, 34, 10, 16, 50, 24, 11, 14, 30], rev_desde=8)

    # --- Reglas combinadas
    filas = []
    for r in KB["reglas_combinadas"]:
        filas.append([r["id"], r.get("titulo", ""), NIVEL.get(r["nivel"], r["nivel"]), condicion(r.get("si")), ", ".join(AMBITOS.get(a, a) for a in r.get("ambitos", [])), r["texto"], "", "", "", f"combinada|{r['id']}"])
    hoja(wb, "Reglas combinadas", ["Id", "Título", "Nivel", "Se activa cuando", "Ámbitos", "Texto actual", "ESTADO", "TEXTO NUEVO", "COMENTARIO"], filas,
         [24, 30, 11, 36, 22, 80, 11, 60, 30], rev_desde=7)

    # --- Interacciones curadas
    filas = []
    for i, x in enumerate(DDI["interacciones"]):
        filas.append([nombre_ref(x["a"]), nombre_ref(x["b"]), x["nivel"], x.get("efecto", ""), x.get("accion", ""), "", "", "", "", f"ddi|{i}"])
    hoja(wb, "Interacciones", ["Fármaco / grupo A", "Fármaco / grupo B", "Nivel", "Efecto", "Qué hacer", "ESTADO", "NIVEL NUEVO (mayor/moderada/menor)", "TEXTO NUEVO (efecto · qué hacer)", "COMENTARIO"], filas,
         [28, 28, 11, 50, 50, 11, 16, 50, 30], rev_desde=6)

    # --- Clases
    filas = []
    for c in KB["clases"]:
        filas.append([c["id"], c["nombre"], c.get("resumen", ""), ", ".join(c.get("senales", [])), ", ".join(c.get("fuentes", [])), "", "", "", f"clase|{c['id']}"])
    hoja(wb, "Clases", ["Id", "Grupo", "Resumen (lo ve el paciente)", "Señales", "Fuentes", "ESTADO", "RESUMEN NUEVO", "COMENTARIO"], filas,
         [22, 34, 70, 30, 30, 11, 60, 30], rev_desde=6)

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    n = sum(ws.max_row - 1 for ws in wb.worksheets[1:])
    print(f"{out} · {n} filas revisables")


if __name__ == "__main__":
    main(Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data" / "revision" / f"revision_{date.today():%Y%m%d}.xlsx")
