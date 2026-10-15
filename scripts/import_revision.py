#!/usr/bin/env python3
"""Aplica una hoja de revisión clínica (generada por export_revision.py y rellenada por el revisor).

Uso:
    python3 scripts/import_revision.py data/revision/revision_20261015.xlsx          # aplica y regenera kb.json
    python3 scripts/import_revision.py revision.xlsx --simular                        # solo muestra qué haría

Lee las columnas ESTADO / TEXTO NUEVO / UMBRAL NUEVO / NIVEL NUEVO / RESUMEN NUEVO y aplica los cambios en las
FUENTES (data/ext/*.json y app/data/kb.json como base; app/data/ddi_curated.json para interacciones), para que
scripts/merge_kb.py no los deshaga. Al final ejecuta merge_kb.py y escribe data/revision/comentarios_<fecha>.md
con todos los comentarios del revisor, para leerlos con calma.
"""
import json
import re
import subprocess
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
KB = ROOT / "app/data/kb.json"
DDI = ROOT / "app/data/ddi_curated.json"
EXT = sorted(p for p in (ROOT / "data/ext").glob("*.json") if not p.name.startswith("ddi_"))
EXT_DDI = sorted((ROOT / "data/ext").glob("ddi_*.json"))
AMB_INV = {"Antes de empezar": "general", "Aeróbico": "aerobico", "Fuerza": "fuerza", "Equilibrio": "equilibrio"}
NIVEL_INV = {"alerta": "alerta", "precaución": "precaucion", "precaucion": "precaucion", "info": "info"}


def cargar(p):
    return json.loads(p.read_text(encoding="utf-8"))


def guardar(p, d):
    p.write_text(json.dumps(d, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def lista_ddi(dd):
    return dd if isinstance(dd, list) else dd.get("interacciones", [])


def celdas(ws):
    cab = [str(c.value or "").strip() for c in ws[1]]
    for fila in ws.iter_rows(min_row=2, values_only=True):
        d = {cab[i]: (str(v).strip() if v is not None else "") for i, v in enumerate(fila) if i < len(cab)}
        if d.get("ref"):
            yield d


def fuente_de(d):
    """Si el comentario empieza por 'Fuente:' o 'FT ...', se guarda como fuente de la regla (p. ej. 'FT 4.4', 'ACSM ap. A')."""
    c = d.get("COMENTARIO", "")
    m = re.match(r"^(?:fuente\s*:\s*)?((?:FT|ACSM|ESC|ADA|Beers|STOPP|Lexicomp|Micromedex|BOT|UpToDate|KDIGO|Vivifrail|GOLD|NICE)[^;.\n]{0,40})", c, re.I)
    return m.group(1).strip() if m else None


def estado(d):
    e = d.get("ESTADO", "").upper()
    nuevo = d.get("TEXTO NUEVO") or d.get("UMBRAL NUEVO") or d.get("NIVEL NUEVO (mayor/moderada/menor)") or d.get("RESUMEN NUEVO") or d.get("TEXTO NUEVO (efecto · qué hacer)")
    if e == "ELIMINAR":
        return "ELIMINAR"
    if e == "CAMBIAR" or (nuevo and e != "OK"):
        return "CAMBIAR"
    return ""


class Fuentes:
    """Todas las fuentes JSON cargadas; cada cambio se aplica en todas las que contengan el elemento."""

    def __init__(self):
        self.docs = {p: cargar(p) for p in EXT}
        self.docs[KB] = cargar(KB)
        self.tocados = set()

    def clases(self):
        for p, d in self.docs.items():
            for c in d.get("clases", []):
                yield p, c, "clase"
            for c in d.get("clases_parche", []):
                yield p, c, "parche"

    def regla(self, cid, amb, texto, nuevo=None, eliminar=False, revisado=None, fuente=None):
        n = 0
        for p, c, tipo in self.clases():
            if c["id"] != cid:
                continue
            lista = c.get(amb if tipo == "clase" else amb + "_add", [])
            for r in list(lista):
                if r["texto"].strip() == texto.strip():
                    if eliminar:
                        lista.remove(r)
                    else:
                        if nuevo:
                            r["texto"] = nuevo
                        if revisado:
                            r["revisado"] = True
                        if fuente:
                            r["fuente"] = fuente
                    self.tocados.add(p)
                    n += 1
        return n

    def principio(self, cid, pid, umbral=None):
        n = 0
        for p, c, tipo in self.clases():
            if c["id"] != cid:
                continue
            for pr in c.get("principios", []) + c.get("principios_add", []):
                if pr["id"] == pid:
                    if umbral is None:
                        pr.pop("umbral_mg", None)
                    else:
                        pr.setdefault("umbral_mg", {})["alta"] = umbral
                    self.tocados.add(p)
                    n += 1
        return n

    def clase(self, cid, resumen):
        n = 0
        for p, c, tipo in self.clases():
            if c["id"] == cid and tipo == "clase":
                c["resumen"] = resumen
                self.tocados.add(p)
                n += 1
        return n

    def combinada(self, rid, nuevo=None, eliminar=False):
        n = 0
        for p, d in self.docs.items():
            lista = d.get("reglas_combinadas", [])
            for r in list(lista):
                if r["id"] == rid:
                    if eliminar:
                        lista.remove(r)
                    else:
                        r["texto"] = nuevo
                    self.tocados.add(p)
                    n += 1
        return n

    def guardar(self):
        for p in self.tocados:
            guardar(p, self.docs[p])
        return len(self.tocados)


def main(xlsx: Path, simular: bool):
    wb = load_workbook(xlsx, data_only=True)
    F = Fuentes()
    ddi = cargar(DDI)
    ddi_ext = {p: cargar(p) for p in EXT_DDI}
    comentarios, log, avisos = [], [], []
    revisadas = [0]

    def nota(d, hoja, contexto):
        if d.get("COMENTARIO"):
            comentarios.append(f"- **{hoja}** · {contexto}\n  {d['COMENTARIO']}")

    # Reglas
    for d in celdas(wb["Reglas"]):
        _, cid, amb, _ = d["ref"].split("|")
        ctx = f"{d['Grupo']} · {d['Ámbito']} · «{d['Texto actual'][:70]}…»"
        nota(d, "Regla", ctx)
        e = estado(d)
        if e == "ELIMINAR":
            n = F.regla(cid, amb, d["Texto actual"], eliminar=True)
            log.append(f"eliminar regla {cid}/{amb} ({n} fuentes): {d['Texto actual'][:60]}…")
        elif e == "CAMBIAR" and d.get("TEXTO NUEVO"):
            n = F.regla(cid, amb, d["Texto actual"], nuevo=d["TEXTO NUEVO"], revisado=True, fuente=fuente_de(d))
            log.append(f"cambiar regla {cid}/{amb} ({n} fuentes): {d['TEXTO NUEVO'][:60]}…")
            if not n:
                avisos.append(f"no encontrada la regla {cid}/{amb}: {d['Texto actual'][:60]}")
        elif e == "CAMBIAR":
            avisos.append(f"CAMBIAR sin TEXTO NUEVO en regla {cid}/{amb} (solo comentario)")
        elif d.get("ESTADO", "").upper() == "OK":
            n = F.regla(cid, amb, d["Texto actual"], revisado=True, fuente=fuente_de(d))
            revisadas[0] += n and 1

    # Umbrales
    for d in celdas(wb["Umbrales"]):
        _, cid, pid = d["ref"].split("|")
        nota(d, "Umbral", f"{d['Principio activo']} (umbral actual {d['Dosis alta a partir de (por día)'] or '—'})")
        e = estado(d)
        if e == "CAMBIAR" and d.get("UMBRAL NUEVO"):
            v = d["UMBRAL NUEVO"].replace(",", ".")
            try:
                v = float(v)
                v = int(v) if v.is_integer() else v
            except ValueError:
                avisos.append(f"umbral no numérico para {pid}: {d['UMBRAL NUEVO']}")
                continue
            n = F.principio(cid, pid, umbral=v)
            log.append(f"umbral {pid} → {v} ({n} fuentes)")
        elif e == "ELIMINAR":
            n = F.principio(cid, pid, umbral=None)
            log.append(f"quitar umbral de {pid} ({n} fuentes)")

    # Reglas combinadas
    for d in celdas(wb["Reglas combinadas"]):
        _, rid = d["ref"].split("|")
        nota(d, "Regla combinada", f"{rid} · {d['Título']}")
        e = estado(d)
        if e == "ELIMINAR":
            log.append(f"eliminar combinada {rid} ({F.combinada(rid, eliminar=True)} fuentes)")
        elif e == "CAMBIAR" and d.get("TEXTO NUEVO"):
            log.append(f"cambiar combinada {rid} ({F.combinada(rid, nuevo=d['TEXTO NUEVO'])} fuentes)")

    # Interacciones curadas
    borrar = []
    for d in celdas(wb["Interacciones"]):
        _, i = d["ref"].split("|")
        i = int(i)
        x = ddi["interacciones"][i] if i < len(ddi["interacciones"]) else None
        ctx = f"{d['Fármaco / grupo A']} + {d['Fármaco / grupo B']}"
        nota(d, "Interacción", ctx)
        if not x:
            continue
        e = estado(d)
        if e == "ELIMINAR":
            borrar.append(i)
            log.append(f"eliminar interacción {ctx}")
        elif e == "CAMBIAR":
            nv = d.get("NIVEL NUEVO (mayor/moderada/menor)", "").lower()
            if nv in ("mayor", "moderada", "menor"):
                x["nivel"] = nv
            tn = d.get("TEXTO NUEVO (efecto · qué hacer)")
            if tn:
                partes = [t.strip() for t in re.split(r"\s·\s|\n", tn, maxsplit=1)]
                x["efecto"] = partes[0]
                if len(partes) > 1:
                    x["accion"] = partes[1]
            log.append(f"cambiar interacción {ctx} → {x['nivel']}")
            # reflejar en el bloque de origen si existe
            for p, dd in ddi_ext.items():
                for y in lista_ddi(dd):
                    if y.get("a") == x["a"] and y.get("b") == x["b"]:
                        y.update({k: x[k] for k in ("nivel", "efecto", "accion") if k in x})
    for i in sorted(borrar, reverse=True):
        x = ddi["interacciones"].pop(i)
        for p, dd in ddi_ext.items():
            lista = lista_ddi(dd)
            for y in list(lista):
                if y.get("a") == x["a"] and y.get("b") == x["b"]:
                    lista.remove(y)

    # Clases
    for d in celdas(wb["Clases"]):
        _, cid = d["ref"].split("|")
        nota(d, "Clase", f"{d['Grupo']}")
        e = estado(d)
        if e == "CAMBIAR" and d.get("RESUMEN NUEVO"):
            log.append(f"resumen de {cid} ({F.clase(cid, d['RESUMEN NUEVO'])} fuentes)")
        elif e == "ELIMINAR":
            avisos.append(f"ELIMINAR clase {cid}: no se aplica automáticamente (hay que quitarla a mano)")

    print(f"Cambios: {len(log)} · reglas marcadas como revisadas (OK): {revisadas[0]} · comentarios: {len(comentarios)} · avisos: {len(avisos)}")
    for l in log:
        print("  ·", l)
    for a in avisos:
        print("  AVISO:", a)
    if simular:
        print("(simulación: no se ha escrito nada)")
        return
    n = F.guardar()
    guardar(DDI, ddi)
    for p, dd in ddi_ext.items():
        guardar(p, dd)
    print(f"Fuentes actualizadas: {n} + ddi_curated.json")
    out = ROOT / "data/revision" / f"comentarios_{date.today():%Y%m%d}.md"
    out.write_text(f"# Comentarios del revisor · {date.today():%d/%m/%Y}\n\n" + ("\n".join(comentarios) if comentarios else "(sin comentarios)") + "\n", encoding="utf-8")
    print(f"Comentarios en {out.relative_to(ROOT)}")
    subprocess.run([sys.executable, str(ROOT / "scripts/merge_kb.py")], check=True)


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not args:
        sys.exit(__doc__)
    main(Path(args[0]), "--simular" in sys.argv)
